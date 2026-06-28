"""PDF Routes"""
import os
import uuid
from flask import Blueprint, render_template, jsonify, request, send_file, redirect, url_for, flash, current_app
from flask_login import login_required, current_user
from werkzeug.utils import secure_filename
from database.models import db, Document
import pdfplumber
from docx import Document as DocxDocument

pdf_bp = Blueprint('pdf', __name__, url_prefix='/pdf')


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in current_app.config['ALLOWED_EXTENSIONS']


def extract_document_text(file_path, file_type):
    if file_type == 'pdf':
        with pdfplumber.open(file_path) as pdf:
            pages = pdf.pages
            text = '\n\n'.join((page.extract_text() or '') for page in pages).strip()
            return text, len(pages)

    if file_type == 'txt':
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read().strip()
            return text, text.count('\n') + 1 if text else 0

    if file_type == 'docx':
        doc = DocxDocument(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        text = '\n\n'.join(paragraphs).strip()
        return text, len(paragraphs)

    return '', 0


@pdf_bp.route('/upload', methods=['GET', 'POST'])
@login_required
def upload():
    if request.method == 'POST':
        file = request.files.get('file')
        title = request.form.get('title', '').strip()
        description = request.form.get('description', '').strip()
        tags = request.form.get('tags', '').strip()

        if not file or file.filename == '':
            flash('Please choose a document file to upload.', 'danger')
            return render_template('upload.html', title=title, description=description, tags=tags)

        if not allowed_file(file.filename):
            flash('Only PDF, TXT, and DOCX files are allowed.', 'danger')
            return render_template('upload.html', title=title, description=description, tags=tags)

        filename = secure_filename(file.filename)
        file_ext = filename.rsplit('.', 1)[1].lower()
        unique_name = f"{uuid.uuid4().hex}.{file_ext}"

        upload_folder = current_app.config['UPLOAD_FOLDER']
        os.makedirs(upload_folder, exist_ok=True)
        file_path = os.path.join(upload_folder, unique_name)
        file.save(file_path)

        document = Document(
            user_id=current_user.id,
            title=title or filename,
            description=description,
            filename=filename,
            file_path=file_path,
            file_size=os.path.getsize(file_path),
            file_type=file_ext,
            mime_type=file.mimetype,
            status='pending',
            tags=tags,
        )

        db.session.add(document)
        db.session.commit()

        try:
            text_content, page_count = extract_document_text(file_path, file_ext)
            document.text_content = text_content
            document.page_count = page_count
            document.status = 'completed'
        except Exception as exc:
            document.status = 'failed'
            document.error_message = str(exc)
        finally:
            db.session.add(document)
            db.session.commit()

        flash('Document uploaded successfully.', 'success')
        return redirect(url_for('pdf.view_pdf', doc_id=document.id))

    return render_template('upload.html')


@pdf_bp.route('/api/upload', methods=['POST'])
@login_required
def api_upload():
    return jsonify({'message': 'Upload endpoint'}), 200


@pdf_bp.route('/<int:doc_id>')
@login_required
def view_pdf(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if doc.user_id != current_user.id:
        return "Unauthorized", 403
    return render_template('pdf_viewer.html', document=doc)


@pdf_bp.route('/<int:doc_id>/file')
@login_required
def serve_file(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if doc.user_id != current_user.id:
        return "Unauthorized", 403
    return send_file(doc.file_path, as_attachment=False, download_name=doc.filename, mimetype=doc.mime_type)


@pdf_bp.route('/<int:doc_id>/download')
@login_required
def download_pdf(doc_id):
    doc = Document.query.get_or_404(doc_id)
    if doc.user_id != current_user.id:
        return "Unauthorized", 403

    return send_file(doc.file_path, as_attachment=True, download_name=doc.filename, mimetype=doc.mime_type)
